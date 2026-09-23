from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "docs" / "数字网格员评委使用说明书.docx"

BLACK = "000000"
NAVY = "17365D"
PALE_BLUE = "EDF3F8"
PALE_GRAY = "F7F7F7"
BORDER = "D9D9D9"
MUTED = RGBColor(89, 89, 89)


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def prevent_row_split(row):
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    tr_pr.append(cant_split)


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=110, start=130, bottom=110, end=130):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_cell_borders(cell, color=BORDER, size="6"):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        node = borders.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            borders.append(node)
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), size)
        node.set(qn("w:color"), color)


def set_run_font(run, east_asia="宋体", latin="Aptos", size=10.5, bold=None, color=BLACK):
    run.font.name = latin
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), east_asia)
    run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)


def set_paragraph_format(paragraph, before=0, after=6, line=1.35):
    fmt = paragraph.paragraph_format
    fmt.space_before = Pt(before)
    fmt.space_after = Pt(after)
    fmt.line_spacing = line
    fmt.widow_control = True


def add_body(doc, text="", bold_lead=None, after=6):
    p = doc.add_paragraph()
    set_paragraph_format(p, after=after)
    if bold_lead and text.startswith(bold_lead):
        lead = p.add_run(bold_lead)
        set_run_font(lead, bold=True)
        rest = p.add_run(text[len(bold_lead):])
        set_run_font(rest)
    else:
        run = p.add_run(text)
        set_run_font(run)
    return p


def add_bullet(doc, text, level=0):
    p = doc.add_paragraph(style="List Bullet" if level == 0 else "List Bullet 2")
    set_paragraph_format(p, after=3, line=1.25)
    for run in p.runs:
        set_run_font(run)
    if not p.runs:
        set_run_font(p.add_run(text))
    else:
        p.runs[0].text = text
    return p


NUMBER_COUNTER = 0


def add_number(doc, text, start=False):
    global NUMBER_COUNTER
    if start:
        NUMBER_COUNTER = 0
    NUMBER_COUNTER += 1
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Cm(0.75)
    p.paragraph_format.first_line_indent = Cm(-0.55)
    set_paragraph_format(p, after=4, line=1.25)
    prefix = p.add_run(f"{NUMBER_COUNTER}. ")
    set_run_font(prefix)
    set_run_font(p.add_run(text))
    return p


def add_code_line(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Cm(0.65)
    p.paragraph_format.right_indent = Cm(0.45)
    set_paragraph_format(p, before=2, after=7, line=1.15)
    run = p.add_run(text)
    set_run_font(run, east_asia="等线", latin="Consolas", size=9.5)
    return p


def add_heading(doc, text, level=1):
    p = doc.add_paragraph(style=f"Heading {level}")
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.page_break_before = False
    p.paragraph_format.space_before = Pt(0 if level == 1 else 10)
    p.paragraph_format.space_after = Pt(7 if level == 1 else 5)
    run = p.add_run(text)
    set_run_font(run, east_asia="微软雅黑", latin="Aptos Display", size=16 if level == 1 else 12, bold=True)
    return p


def add_table(doc, headers, rows, widths=None, center_cols=None, font_size=9.2):
    center_cols = set(center_cols or [])
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    hdr = table.rows[0]
    set_repeat_table_header(hdr)
    for idx, title in enumerate(headers):
        cell = hdr.cells[idx]
        if widths:
            cell.width = Cm(widths[idx])
        set_cell_shading(cell, NAVY)
        set_cell_margins(cell)
        set_cell_borders(cell)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        set_paragraph_format(p, after=0, line=1.1)
        run = p.add_run(title)
        set_run_font(run, east_asia="微软雅黑", size=9.2, bold=True, color="FFFFFF")
    for row_idx, values in enumerate(rows):
        row = table.add_row()
        prevent_row_split(row)
        for col_idx, value in enumerate(values):
            cell = row.cells[col_idx]
            if widths:
                cell.width = Cm(widths[col_idx])
            set_cell_shading(cell, PALE_BLUE if row_idx % 2 else "FFFFFF")
            set_cell_margins(cell)
            set_cell_borders(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if col_idx in center_cols else WD_ALIGN_PARAGRAPH.LEFT
            set_paragraph_format(p, after=0, line=1.2)
            run = p.add_run(str(value))
            set_run_font(run, size=font_size)
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(3)
    return table


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    set_paragraph_format(paragraph, after=0, line=1)
    prefix = paragraph.add_run("数字网格员评委使用说明书    第 ")
    set_run_font(prefix, east_asia="微软雅黑", size=8.5, color="666666")
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = "PAGE"
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run = paragraph.add_run()
    run._r.append(fld_char1)
    run._r.append(instr)
    run._r.append(fld_char2)
    suffix = paragraph.add_run(" 页")
    set_run_font(suffix, east_asia="微软雅黑", size=8.5, color="666666")


def build():
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Cm(2.1)
    section.bottom_margin = Cm(1.8)
    section.left_margin = Cm(2.35)
    section.right_margin = Cm(2.35)
    section.header_distance = Cm(0.9)
    section.footer_distance = Cm(0.8)

    normal = doc.styles["Normal"]
    normal.font.name = "Aptos"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor.from_string(BLACK)

    for style_name, size in (("Title", 24), ("Subtitle", 12), ("Heading 1", 16), ("Heading 2", 12)):
        style = doc.styles[style_name]
        style.font.name = "Aptos Display"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor.from_string(BLACK)
        style.font.bold = style_name != "Subtitle"

    title_ppr = doc.styles["Title"]._element.get_or_add_pPr()
    title_border = title_ppr.find(qn("w:pBdr"))
    if title_border is not None:
        title_ppr.remove(title_border)

    props = doc.core_properties
    props.title = "数字网格员评委使用说明书"
    props.subject = "数字网格员参赛作品启动体验与验收指南"
    props.author = "玄黄二人组"
    props.keywords = "数字网格员, 智慧政务, 社区治理, HiAgent, 微信小程序"

    # Cover
    p = doc.add_paragraph(style="Title")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(92)
    p.paragraph_format.space_after = Pt(16)
    run = p.add_run("数字网格员评委使用说明书")
    set_run_font(run, east_asia="微软雅黑", latin="Aptos Display", size=24, bold=True)

    sub = doc.add_paragraph(style="Subtitle")
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub.paragraph_format.space_after = Pt(54)
    run = sub.add_run("Windows 一键启动  网页端  微信小程序  评审演示")
    set_run_font(run, east_asia="微软雅黑", size=12, color="404040")

    intro = doc.add_paragraph()
    intro.alignment = WD_ALIGN_PARAGRAPH.CENTER
    intro.paragraph_format.left_indent = Cm(1.4)
    intro.paragraph_format.right_indent = Cm(1.4)
    set_paragraph_format(intro, after=24, line=1.55)
    run = intro.add_run("本说明书帮助评委在不修改源码、不配置真实密钥的情况下，独立完成系统启动、核心流程体验和功能验收。建议先使用本地模拟模式完成评审，再按需验证 HiAgent 真实接口。")
    set_run_font(run, size=11)

    meta = doc.add_paragraph()
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_paragraph_format(meta, before=34, after=0, line=1.5)
    r = meta.add_run("参赛队伍  玄黄二人组\n参赛方向  智慧政务与社区网格治理\n文档版本  1.0\n更新日期  2026 年 9 月 13 日")
    set_run_font(r, east_asia="微软雅黑", size=10, color="595959")

    add_page_number(section.footer.paragraphs[0])
    doc.add_page_break()

    # Opening and quick start
    add_heading(doc, "一 评审结论与体验范围", 1)
    add_body(doc, "数字网格员是一套面向社区治理的 AI Agent 协同系统。评委可通过网页端与微信小程序体验居民登记、诉求受理、意图识别、政策检索、工单生成、自动分派、权限隔离和人工确认归档。系统默认提供本地模拟模式，因此没有 HiAgent API Key 也能完成完整演示。")
    add_body(doc, "一键启动的范围：启动程序负责准备 Node.js 运行环境、启动本地后端、打开网页端，并尝试通过微信开发者工具 CLI 导入小程序。它用于本地评审和联调，不等同于将小程序自动提交微信审核或发布到生产环境。", bold_lead="一键启动的范围：")

    add_heading(doc, "二 三分钟快速启动", 1)
    add_heading(doc, "2.1 评审前准备", 2)
    add_table(
        doc,
        ["项目", "最低要求", "说明"],
        [
            ["操作系统", "Windows 10 或 Windows 11", "项目目录需要具备写入权限"],
            ["网页体验", "任意现代浏览器", "无需安装额外前端依赖"],
            ["小程序体验", "微信开发者工具", "首次使用需扫码登录并开启服务端口"],
            ["网络", "本地模拟可离线运行", "自动下载 Node.js 或真实接口模式需要联网"],
        ],
        widths=[3.2, 5.0, 7.0],
        center_cols=[0, 1],
    )
    add_heading(doc, "2.2 启动步骤", 2)
    add_number(doc, "完整解压项目，进入 数字网格员_Windows一键启动_UTF8\\数字网格员 目录。", start=True)
    add_number(doc, "双击 启动数字网格员.bat。请勿运行旧版 ASCII 或 WindowsOneClick 副本。")
    add_number(doc, "首次出现 是否配置 HiAgent 真实接口 提示时，评审演示直接按回车。")
    add_number(doc, "等待网页端和微信开发者工具自动打开。若防火墙询问是否允许访问，请选择允许。")
    add_number(doc, "使用启动窗口显示的账号登录网页端，开始体验。")
    add_body(doc, "正常启动标志：窗口显示“数字网格员已启动”，网页地址通常为 http://127.0.0.1:3100，微信开发者工具显示项目 digital-grid-worker-miniprogram。", bold_lead="正常启动标志：")

    add_heading(doc, "2.3 测试账号", 2)
    add_table(
        doc,
        ["角色", "账号", "密码", "可见范围"],
        [
            ["系统管理员", "admin", "admin123", "全部工单和全部网格员"],
            ["华林社区负责人", "hualin_admin", "123456", "华林社区工单"],
            ["华林网格员 01", "hualin01", "123456", "派给本人的华林网格 01 工单"],
            ["华林网格员 02", "hualin02", "123456", "派给本人的华林网格 02 工单"],
            ["观风亭网格员 01", "guanfeng01", "123456", "派给本人的观风亭网格 01 工单"],
        ],
        widths=[3.5, 3.2, 2.7, 6.0],
        center_cols=[0, 1, 2],
    )

    # Web app
    add_heading(doc, "三 网页端操作说明", 1)
    add_heading(doc, "3.1 登录和导航", 2)
    add_body(doc, "网页未自动打开时，在浏览器访问启动窗口显示的本地地址。使用管理员账号登录后，左侧导航提供诉求受理台、工单中心、知识库、数据看板和系统配置。评审建议先使用管理员账号完整体验，再切换普通网格员账号验证权限隔离。")

    add_heading(doc, "3.2 诉求受理与 Agent 链路", 2)
    add_number(doc, "进入诉求受理台，选择已有居民档案或新增居民登记。", start=True)
    add_number(doc, "输入居民诉求，或直接使用系统内置示例场景。")
    add_number(doc, "提交后观察右侧处理链路，核对意图类别、责任部门、响应时限和政策引用。")
    add_number(doc, "投诉类事项需补齐地址、时间、频率等要素；居民确认后才生成正式工单。")
    add_number(doc, "处理完成后进入工单中心，由网格员确认处置和归档。")

    add_heading(doc, "3.3 居民档案与 Excel 导入", 2)
    add_body(doc, "居民档案支持姓名、手机号和楼栋房号检索，可新增登记或修改地址。历史工单继续关联原居民身份，地址变化会作为档案变更保留。")
    add_bullet(doc, "单个登记：适合现场代办、老年居民或临时补录。")
    add_bullet(doc, "批量导入：下载 data\\templates\\居民信息登记表.xlsx，回收后导入 .xlsx 或 .csv。")
    add_bullet(doc, "自动识别：按手机号或姓名判断新增、更新和重复记录，缺少姓名或地址的行会被跳过。")

    add_heading(doc, "3.4 工单与权限验证", 2)
    add_body(doc, "管理员和社区负责人可以查看权限范围内的工单并改派责任网格员。普通网格员只看到派给本人且属于本人网格的工单。权限判断由后端执行，不依赖前端隐藏。")
    add_number(doc, "使用 admin 登录，记录一条工单的编号和当前责任网格员。", start=True)
    add_number(doc, "将该工单改派给 hualin01，然后退出管理员账号。")
    add_number(doc, "使用 hualin01 登录，确认能够看到该工单。")
    add_number(doc, "再使用 hualin02 登录，确认不能看到未派给本人的工单。")

    # Mini program
    add_heading(doc, "四 微信小程序操作说明", 1)
    add_heading(doc, "4.1 开发者工具模拟器", 2)
    add_body(doc, "一键启动会调用微信开发者工具打开项目根目录，project.config.json 已指定 miniprogram 作为小程序目录。首次调用 CLI 前，应在微信开发者工具中完成扫码登录，并在设置的安全设置中开启服务端口。")
    add_number(doc, "确认开发者工具显示 AppID wxdc9dee66a166b036。", start=True)
    add_number(doc, "点击编译，等待首页出现。")
    add_number(doc, "首次进入时完成居民登记，然后进入对话页面提交诉求。")
    add_number(doc, "在工单页面查看本人提交的工单和处理状态。")

    add_heading(doc, "4.2 真机预览", 2)
    add_body(doc, "手机中的 127.0.0.1 指向手机本身，不能用于访问评审电脑。真机测试时，手机和电脑必须连接同一局域网，并在小程序“我的”页面将后端 API 地址改为启动窗口显示的电脑局域网地址。")
    add_code_line(doc, "示例  http://192.168.1.100:3100")
    add_body(doc, "如果电脑同时安装 WSL、VMware、ZeroTier 等虚拟网卡，应优先选择与手机同一网段的物理 Wi-Fi 或以太网地址。正式上线必须使用公网 HTTPS，并在微信公众平台配置 request 合法域名。")

    # Demo workflow
    doc.add_page_break()
    add_heading(doc, "五 推荐的五分钟评审演示", 1)
    add_table(
        doc,
        ["时间", "操作", "评审观察点"],
        [
            ["0:00 至 0:40", "双击启动并使用 admin 登录", "零依赖启动、本地模拟模式、网页与小程序同时就绪"],
            ["0:40 至 1:20", "检索居民并修改一次地址", "档案检索、地址变更、历史关联不丢失"],
            ["1:20 至 2:20", "提交环境卫生诉求", "意图识别、要素抽取、自动派单、响应时限"],
            ["2:20 至 3:20", "提交噪音或政策咨询", "RAG 检索、政策引用、来源可追溯"],
            ["3:20 至 4:20", "查看工单并改派网格员", "结构化工单、权限范围、人工确认"],
            ["4:20 至 5:00", "切换普通账号并打开看板", "权限隔离、闭环记录、指标汇总"],
        ],
        widths=[3.1, 5.3, 7.0],
        center_cols=[0],
    )

    add_heading(doc, "5.1 推荐测试语句", 2)
    add_body(doc, "环境卫生场景")
    add_code_line(doc, "楼下垃圾堆了好几天没人清，天气热味道很大，地址在华林花园 3 号楼门口，麻烦尽快处理。")
    add_body(doc, "预期结果：识别为环境卫生类，责任部门为鼓楼区环卫中心，响应时限为 24 小时，并生成包含事项、地址和责任人的工单。")
    add_body(doc, "噪音扰民场景")
    add_code_line(doc, "五四路 193 号旁边的工地半夜还在施工，声音很大，吵得睡不着，请问这种情况怎么投诉？")
    add_body(doc, "预期结果：识别为噪音扰民类，关联城管处置，并引用噪声污染防治相关政策依据。")

    # Architecture and differentiation
    add_heading(doc, "六 技术架构与评审要点", 1)
    add_table(
        doc,
        ["处理阶段", "系统动作", "可验证输出"],
        [
            ["诉求接入", "统一接收网页、小程序和模拟渠道输入", "居民身份、原始文本、事发地址"],
            ["意图识别", "判断类别并提取时间、地点、频率、紧急度", "分类结果和结构化要素"],
            ["上下文补全", "关联居民档案、历史诉求和网格归属", "补全后的请求上下文"],
            ["知识检索", "从 21 篇官方政策知识中检索相关内容", "政策标题、文号、原文链接"],
            ["工单生成", "匹配责任部门、网格员、优先级和时限", "标准化工单"],
            ["人工确认", "由居民确认提交，由网格员确认归档", "可追踪的闭环状态"],
        ],
        widths=[3.1, 6.4, 5.9],
        center_cols=[0],
    )
    add_body(doc, "多 Agent 可拆分升级：意图识别、政策解答和工单处理采用清晰的职责边界，任一节点均可独立优化。")
    add_body(doc, "RAG 可溯源：知识库包含 21 篇福州市、福建省及鼓楼区公开政策资料，覆盖市容环卫、物业、违建、养老、消防、养犬、惠企等类别。")
    add_body(doc, "人机协同可控：Agent 负责高频标准化环节，低置信度和高风险事项可转人工，归档动作必须由网格员确认。")
    add_body(doc, "权限由后端保证：管理员、社区负责人和普通网格员拥有不同可见范围，避免仅靠页面隐藏造成越权。")

    # Security
    add_heading(doc, "七 数据与安全说明", 1)
    add_bullet(doc, "HiAgent API Key 仅保存在后端运行目录，不写入网页公开配置或微信小程序。")
    add_bullet(doc, "本地演示数据保存在 data\\runtime\\store.json，适用于评审和原型验证。")
    add_bullet(doc, "居民姓名、手机号和住址属于个人信息，批量登记表应通过受控渠道收集并及时清理。")
    add_bullet(doc, "正式上线前应将 JSON 存储替换为正式数据库，并补充备份、审计、密钥托管和 HTTPS。")
    add_bullet(doc, "政策资料保留标题、文号、发布日期和官方链接；正式业务使用前仍应进行时效性校对。")

    # Troubleshooting
    doc.add_page_break()
    add_heading(doc, "八 常见问题与处理", 1)
    add_table(
        doc,
        ["现象", "可能原因", "处理方法"],
        [
            ["出现 Profile is not recognized", "运行了旧版 UTF-8 或 LF 换行的批处理副本", "只运行当前 UTF8 包内的 启动数字网格员.bat"],
            ["窗口提示配置 HiAgent", "首次启动尚无真实接口配置", "评审演示直接按回车，进入本地模拟模式"],
            ["网页打不开", "后端尚未就绪或端口被占用", "查看启动窗口给出的实际端口；脚本会自动尝试下一端口"],
            ["开发者工具未自动打开", "工具未安装、未登录或服务端口关闭", "启动开发者工具，扫码登录，并在安全设置中开启服务端口"],
            ["模拟器请求失败", "小程序地址和后端端口不一致", "重新运行启动程序，确认 config.js 中端口与启动窗口一致"],
            ["真机请求超时", "使用了 127.0.0.1 或选中了虚拟网卡地址", "在“我的”页面填写电脑物理网卡地址，确保手机与电脑同网段"],
            ["真实 HiAgent 调用失败", "应用未发布或接口参数不完整", "核对 Endpoint、应用 ID、工作流 ID、API Key 和发布状态"],
        ],
        widths=[4.4, 5.0, 6.0],
        center_cols=[],
        font_size=8.8,
    )
    add_heading(doc, "8.1 停止和重新启动", 2)
    add_body(doc, "启动程序会在独立 Node.js 窗口中运行后端。关闭该 Node.js 窗口即可停止服务。再次启动前，如果原服务仍在运行，脚本会选择下一个空闲端口，因此应以新启动窗口显示的地址为准。")

    # Checklist
    add_heading(doc, "九 评审验收清单", 1)
    add_table(
        doc,
        ["序号", "验收项目", "通过标准", "结果"],
        [
            ["1", "一键启动", "后端健康检查通过，网页端和开发者工具正常打开", "□"],
            ["2", "居民档案", "可新增、检索、编辑并保留地址变更", "□"],
            ["3", "诉求识别", "能识别常见类别并提取地址、时间等信息", "□"],
            ["4", "知识检索", "政策回答显示引用来源，可进行人工复核", "□"],
            ["5", "工单闭环", "可生成、分派、查看进度并由网格员确认归档", "□"],
            ["6", "权限隔离", "不同角色只能查看各自权限范围内的数据", "□"],
            ["7", "小程序", "可登记、对话并查看本人工单", "□"],
            ["8", "离线演示", "没有 HiAgent API Key 时仍能完成核心流程", "□"],
        ],
        widths=[1.5, 3.6, 8.8, 1.5],
        center_cols=[0, 1, 3],
        font_size=9.0,
    )

    add_heading(doc, "十 文件与目录速查", 1)
    add_table(
        doc,
        ["文件或目录", "用途"],
        [
            ["启动数字网格员.bat", "Windows 评审环境首选双击入口"],
            ["一键启动数字网格员.ps1", "环境检测、端口选择、后端启动和开发者工具调用"],
            ["server.js", "本地静态服务、身份认证、工单与 HiAgent 代理"],
            ["index.html", "数字网格员网页端入口"],
            ["miniprogram", "居民端微信小程序代码"],
            ["data\\hiagent-import", "可导入 HiAgent 的 21 篇政策资料"],
            ["data\\templates", "居民批量登记 Excel 模板"],
            ["docs", "方案、演示材料和使用说明"],
        ],
        widths=[6.0, 9.4],
        center_cols=[0],
    )

    add_body(doc, "评审建议：先用本地模拟模式验证产品闭环和交互，再验证真实 HiAgent 配置。这样可以将本机环境问题与智能体接口问题分开，快速判断作品本身的完整性。", bold_lead="评审建议：")

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build()
