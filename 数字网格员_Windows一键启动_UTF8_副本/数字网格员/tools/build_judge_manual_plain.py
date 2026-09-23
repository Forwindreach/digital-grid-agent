from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.enum.section import WD_SECTION_START
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "docs" / "数字网格员评委操作说明书通俗版.docx"

BLACK = "000000"
GRAY = "666666"
LIGHT_GRAY = "D9D9D9"
PALE = "EFF7F4"
GREEN = "0B6B4F"
WHITE = "FFFFFF"


def set_run_font(run, size=10.5, bold=False, color=BLACK, east_asia="微软雅黑"):
    run.font.name = "Aptos"
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), east_asia)
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), "Aptos")
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), "Aptos")
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_borders(cell, color=LIGHT_GRAY, size="6"):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = "w:" + edge
        node = borders.find(qn(tag))
        if node is None:
            node = OxmlElement(tag)
            borders.append(node)
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), size)
        node.set(qn("w:color"), color)


def set_cell_margins(cell, top=120, start=150, bottom=120, end=150):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, v in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn("w:" + m))
        if node is None:
            node = OxmlElement("w:" + m)
            tc_mar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def keep_with_next(paragraph):
    paragraph.paragraph_format.keep_with_next = True


def add_title(doc, text):
    p = doc.add_paragraph(style="Title")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(80)
    p.paragraph_format.space_after = Pt(16)
    r = p.add_run(text)
    set_run_font(r, size=24, bold=True)
    return p


def add_heading(doc, text, level=1):
    p = doc.add_paragraph(style=f"Heading {level}")
    p.paragraph_format.space_before = Pt(12 if level == 1 else 8)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.keep_with_next = True
    r = p.add_run(text)
    set_run_font(r, size=16 if level == 1 else 12, bold=True)
    return p


def add_body(doc, text, bold_lead=None):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.35
    if bold_lead and text.startswith(bold_lead):
        r1 = p.add_run(bold_lead)
        set_run_font(r1, bold=True)
        r2 = p.add_run(text[len(bold_lead):])
        set_run_font(r2)
    else:
        r = p.add_run(text)
        set_run_font(r)
    return p


def add_bullet(doc, text):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.25
    r = p.add_run(text)
    set_run_font(r)
    return p


def add_check(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Cm(0.2)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.0
    r = p.add_run("□ " + text)
    set_run_font(r, size=10.2)
    return p


def add_number(doc, text, start=False):
    p = doc.add_paragraph(style="List Number")
    if start:
        num_pr = p._p.get_or_add_pPr().get_or_add_numPr()
        ilvl = num_pr.get_or_add_ilvl()
        ilvl.val = 0
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.3
    r = p.add_run(text)
    set_run_font(r, size=10.8)
    return p


def add_step(doc, number, title, text, expected=None):
    table = doc.add_table(rows=1, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    table.columns[0].width = Cm(1.25)
    table.columns[1].width = Cm(14.3)
    left, right = table.rows[0].cells
    left.width = Cm(1.25)
    right.width = Cm(14.3)
    set_cell_shading(left, GREEN)
    set_cell_shading(right, WHITE)
    for cell in (left, right):
        set_cell_borders(cell)
        set_cell_margins(cell, top=150, start=160, bottom=150, end=160)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    p = left.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(0)
    r = p.add_run(str(number))
    set_run_font(r, size=16, bold=True, color=WHITE)
    p = right.paragraphs[0]
    p.paragraph_format.space_after = Pt(3)
    r = p.add_run(title)
    set_run_font(r, size=11.5, bold=True)
    p = right.add_paragraph()
    p.paragraph_format.space_after = Pt(2 if expected else 0)
    p.paragraph_format.line_spacing = 1.25
    r = p.add_run(text)
    set_run_font(r, size=10.3)
    if expected:
        p = right.add_paragraph()
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = 1.2
        r = p.add_run("看到的结果  ")
        set_run_font(r, size=9.6, bold=True, color=GREEN)
        r = p.add_run(expected)
        set_run_font(r, size=9.6, color=GRAY)
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(2)


def add_table(doc, headers, rows, widths=None, center_cols=()):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    tr_pr = table.rows[0]._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)
    for idx, header in enumerate(headers):
        cell = table.rows[0].cells[idx]
        if widths:
            cell.width = Cm(widths[idx])
        set_cell_shading(cell, GREEN)
        set_cell_borders(cell)
        set_cell_margins(cell)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run(header)
        set_run_font(r, size=9.6, bold=True, color=WHITE)
    for row_idx, values in enumerate(rows):
        cells = table.add_row().cells
        for idx, value in enumerate(values):
            cell = cells[idx]
            if widths:
                cell.width = Cm(widths[idx])
            set_cell_shading(cell, PALE if row_idx % 2 else WHITE)
            set_cell_borders(cell)
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if idx in center_cols else WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.2
            r = p.add_run(str(value))
            set_run_font(r, size=9.4)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return table


def add_page_number(section):
    p = section.footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(0)
    r = p.add_run("数字网格员评委操作说明书   ")
    set_run_font(r, size=8.5, color=GRAY)
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = "PAGE"
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    run = p.add_run()
    run._r.append(fld_begin)
    run._r.append(instr)
    run._r.append(fld_end)


def page_break(doc):
    doc.add_section(WD_SECTION_START.NEW_PAGE)


def build():
    doc = Document()
    section = doc.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(1.9)
    section.bottom_margin = Cm(1.7)
    section.left_margin = Cm(2.35)
    section.right_margin = Cm(2.35)
    section.header_distance = Cm(0.8)
    section.footer_distance = Cm(0.75)

    normal = doc.styles["Normal"]
    normal.font.name = "Aptos"
    normal._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "微软雅黑")
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor.from_string(BLACK)

    for style_name, size, bold in (
        ("Title", 24, True),
        ("Subtitle", 12, False),
        ("Heading 1", 16, True),
        ("Heading 2", 12, True),
    ):
        style = doc.styles[style_name]
        style.font.name = "Aptos"
        style._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "微软雅黑")
        style.font.size = Pt(size)
        style.font.bold = bold
        style.font.color.rgb = RGBColor.from_string(BLACK)

    title_ppr = doc.styles["Title"]._element.get_or_add_pPr()
    border = title_ppr.find(qn("w:pBdr"))
    if border is not None:
        title_ppr.remove(border)

    props = doc.core_properties
    props.title = "数字网格员评委操作说明书"
    props.subject = "非计算机专业评委使用指南"
    props.author = "玄黄二人组"
    props.keywords = "数字网格员, 评委, 一键启动, 微信小程序, 真机调试"

    add_page_number(section)

    # Cover and shortest path
    add_title(doc, "数字网格员评委操作说明书")
    p = doc.add_paragraph(style="Subtitle")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(34)
    r = p.add_run("非计算机专业评委适用")
    set_run_font(r, size=12, color=GRAY)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.left_indent = Cm(1.4)
    p.paragraph_format.right_indent = Cm(1.4)
    p.paragraph_format.space_after = Pt(24)
    p.paragraph_format.line_spacing = 1.5
    r = p.add_run("您不需要修改代码，也不需要配置任何密钥。完整解压文件后，双击启动文件即可先体验网页端；需要验证微信小程序时，再按本说明书使用自己的微信扫码。")
    set_run_font(r, size=11)

    add_heading(doc, "最短操作路线", 1)
    add_table(
        doc,
        ["顺序", "您要做的事", "完成标志"],
        [
            ["1", "完整解压项目压缩包", "能看到 数字网格员 文件夹"],
            ["2", "双击 启动数字网格员.bat", "出现黑色启动窗口"],
            ["3", "浏览器登录网页端", "进入诉求受理台"],
            ["4", "用自己的微信登录开发者工具", "看到数字网格员首页"],
            ["5", "需要时点击真机调试并用手机扫码", "手机中打开小程序"],
        ],
        widths=[1.8, 7.0, 6.7],
        center_cols=(0,),
    )
    add_body(doc, "只体验网页端：完成前三步即可，不需要微信开发者权限，也不需要手机扫码。", bold_lead="只体验网页端：")
    add_body(doc, "完整体验小程序：参赛方需要提前把您的微信号添加为该小程序的开发者。您始终使用自己的微信，不需要借用参赛者账号。", bold_lead="完整体验小程序：")

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(28)
    r = p.add_run("参赛队伍  玄黄二人组\n文档版本  2.0\n更新日期  2026 年 9 月 13 日")
    set_run_font(r, size=9.8, color=GRAY)

    # Preparation and startup
    page_break(doc)
    add_heading(doc, "开始前准备", 1)
    add_table(
        doc,
        ["准备项目", "要求"],
        [
            ["电脑", "Windows 10 或 Windows 11"],
            ["网页体验", "Edge 或 Chrome 浏览器"],
            ["小程序体验", "已经安装微信开发者工具"],
            ["手机真机体验", "手机微信与电脑连接同一个 Wi Fi"],
            ["微信权限", "参赛方已把评委微信号添加为开发者"],
        ],
        widths=[4.8, 10.7],
        center_cols=(0,),
    )
    add_body(doc, "请勿直接在压缩包中双击启动文件。应先选择“全部解压”，再进入解压后的 数字网格员 文件夹。")

    add_heading(doc, "一键启动", 1)
    add_step(doc, 1, "找到启动文件", "打开解压后的 数字网格员 文件夹，找到 启动数字网格员.bat。文件图标通常是黑色窗口。")
    add_step(doc, 2, "双击运行", "双击 启动数字网格员.bat。若 Windows 显示保护提示，可点击 更多信息，再点击 仍要运行。", "出现黑色窗口，并开始检查运行环境。")
    add_step(doc, 3, "选择本地演示", "如果窗口询问“是否配置 HiAgent 真实接口”，直接按键盘上的 Enter 键。评审不需要输入 API Key。", "窗口继续运行，不再等待输入。")
    add_step(doc, 4, "允许防火墙访问", "若 Windows 防火墙弹出提示，勾选专用网络并点击 允许访问。手机真机调试需要这项权限。")
    add_step(doc, 5, "等待程序打开", "首次启动可能需要下载运行组件，请保持网络连接并等待。不要关闭黑色窗口。", "浏览器打开登录页，微信开发者工具也会尝试自动打开。")
    add_body(doc, "请保持启动窗口开启。关闭运行后端的黑色窗口后，网页和手机小程序将无法继续访问数据。", bold_lead="请保持启动窗口开启。")

    # Web experience
    add_heading(doc, "网页端体验", 1)
    add_body(doc, "网页端是最容易完成的评审入口。即使微信小程序暂时无法扫码，网页端仍可展示主要功能。")
    add_step(doc, 1, "登录系统", "网页自动打开后，输入账号 admin，密码 admin123，然后点击 登录。", "左侧出现诉求受理台 工单中心 政策知识库 数据看板等菜单。")
    add_step(doc, 2, "提交一条居民诉求", "在诉求受理台选择 环境卫生 示例，然后点击 提交并启动 Agent 处理。", "右侧依次显示诉求识别 信息补全 处置参考 工单生成和居民回复。")
    add_step(doc, 3, "查看工单", "点击左侧 工单中心，找到刚生成的记录并打开。", "可以查看工单编号 居民信息 责任部门 处理状态和时间。")
    add_step(doc, 4, "查看政策来源", "点击左侧 政策知识库，选择一条政策资料。", "页面显示政策标题 文号 发布日期和原文链接。")
    add_step(doc, 5, "查看数据结果", "点击左侧 数据看板。", "页面显示受理数量 响应时间 自动闭环率等指标。")

    add_heading(doc, "建议测试内容", 2)
    add_table(
        doc,
        ["测试类型", "可直接使用的示例"],
        [
            ["环境卫生", "楼下垃圾堆了几天，地址在华林花园三号楼门口，请尽快处理。"],
            ["噪音扰民", "五四路附近的工地半夜施工，声音很大，应该向哪里反映？"],
            ["政策咨询", "福州办理养犬登记需要准备哪些材料？"],
        ],
        widths=[3.3, 12.2],
        center_cols=(0,),
    )
    add_body(doc, "评审演示数据均为模拟数据。请勿输入真实居民的姓名 电话或住址。", bold_lead="评审演示数据均为模拟数据。")

    # DevTools login
    add_heading(doc, "微信开发者工具登录", 1)
    add_body(doc, "这一部分用于验证居民端微信小程序。评委应使用自己的微信登录。添加开发者权限只是允许该微信打开项目，不会要求评委提供微信密码。")
    add_step(doc, 1, "确认开发者工具已打开", "一键启动通常会自动打开微信开发者工具。如果没有打开，请手动启动微信开发者工具，再选择 导入项目，并选择整个 数字网格员 文件夹。")
    add_step(doc, 2, "用评委自己的微信扫码", "开发者工具显示登录二维码后，用评委本人手机微信扫一扫并确认登录。不要使用参赛者微信。", "电脑端显示评委自己的微信头像。")
    add_step(doc, 3, "保持项目原有 AppID", "项目已经预先填写 AppID，不要改成测试号，也不要点击无 AppID。如果提示无权限，请联系参赛方确认添加的微信号是否与当前登录微信一致。")
    add_step(doc, 4, "点击编译", "点击开发者工具上方的 编译。首次编译请等待页面加载完成。", "中间的手机模拟器显示 数字网格员 首页。")
    add_step(doc, 5, "先在模拟器体验", "在模拟器中点击 登记，填写模拟信息并保存；随后点击 对话 输入诉求。", "可以收到数字网格员回复，并在 工单 页面看到记录。")

    add_heading(doc, "没有开发者权限时", 2)
    add_body(doc, "如果页面提示“当前微信没有开发权限”或“无权限访问此 AppID”，评委无需反复扫码。请把当前微信号发给参赛方，由参赛方在小程序后台添加开发者。权限生效后，退出开发者工具并重新登录。")

    # Real device
    add_heading(doc, "手机真机调试", 1)
    add_body(doc, "安卓手机和苹果手机的操作相同。手机只是运行小程序，数据服务仍在评委电脑上，因此手机与电脑必须处于同一个网络。")
    add_step(doc, 1, "确认网络", "让电脑和手机连接同一个 Wi Fi，并暂时关闭电脑或手机上的 VPN 代理。")
    add_step(doc, 2, "记下电脑地址", "查看黑色启动窗口，找到“小程序真机后端”后面的地址，例如 http://192.168.1.100:3100。实际数字以窗口显示为准。")
    add_step(doc, 3, "生成真机二维码", "在微信开发者工具上方点击 真机调试。出现选择项时使用默认模式即可。", "电脑屏幕显示一个二维码。")
    add_step(doc, 4, "用同一微信扫码", "使用刚才登录开发者工具的同一个微信扫描二维码，并在手机上确认。", "手机打开 数字网格员。")
    add_step(doc, 5, "检查后端地址", "在小程序底部点击 我的，查看 后端 API 地址。若无法连接，把它改成第 2 步记下的地址，再点击 测试连接 和 保存并重新登录。", "出现连接成功提示。")
    add_step(doc, 6, "完成居民端流程", "使用模拟姓名 电话和住址完成登记，点击 对话 描述问题。若系统请您确认提交，回复“确认提交”，再到 工单 页面查看结果。")

    add_heading(doc, "手机无法连接时", 2)
    add_bullet(doc, "确认黑色启动窗口仍然开着。")
    add_bullet(doc, "确认手机和电脑连接的是同一个 Wi Fi。")
    add_bullet(doc, "确认手机填写的不是 127.0.0.1，而是启动窗口显示的小程序真机后端地址。")
    add_bullet(doc, "若防火墙曾被拒绝，请重新启动程序并允许专用网络访问。")
    add_bullet(doc, "仍无法连接时，可以先完成网页端评审，不影响主要功能展示。")

    # Recommended test and explanations
    add_heading(doc, "五分钟推荐体验", 1)
    add_table(
        doc,
        ["用时", "操作", "重点观察"],
        [
            ["一分钟", "双击启动并登录网页端", "是否能一键打开系统"],
            ["一分钟", "提交环境卫生示例", "是否自动识别问题并生成处理步骤"],
            ["一分钟", "打开工单中心和政策知识库", "工单是否完整 政策是否可追溯"],
            ["一分钟", "用小程序登记并提交诉求", "居民端是否能收到回复"],
            ["一分钟", "打开我的工单和数据看板", "处理过程是否形成闭环记录"],
        ],
        widths=[2.4, 6.4, 6.7],
        center_cols=(0,),
    )

    add_heading(doc, "常见词语解释", 1)
    add_table(
        doc,
        ["页面中的词", "简单解释"],
        [
            ["后端", "在电脑上运行的数据服务。关闭黑色窗口后，它就停止工作。"],
            ["AppID", "微信为每个小程序分配的编号。项目已经填好，评委无需修改。"],
            ["编译", "让开发者工具重新读取项目并显示最新页面。"],
            ["真机调试", "让小程序临时运行在评委自己的手机上。"],
            ["本地模拟模式", "不连接真实 AI 平台也能演示完整流程的安全评审模式。"],
            ["HiAgent", "项目可接入的智能体平台。普通评审不需要填写密钥。"],
        ],
        widths=[4.0, 11.5],
        center_cols=(0,),
    )
    add_body(doc, "本说明书中的真机调试仅用于现场评审，不代表小程序已经正式发布。正式上线还需要完成微信平台审核和服务器配置。")

    # Troubleshooting and finish
    add_heading(doc, "常见问题处理", 1)
    add_table(
        doc,
        ["看到的情况", "请这样处理"],
        [
            ["提示 Profile is not recognized", "关闭窗口，确认运行的是当前文件夹中的 启动数字网格员.bat，不要运行旧副本。"],
            ["窗口一直停在 HiAgent 配置", "直接按 Enter 键，不要输入密钥。"],
            ["网页没有自动打开", "在黑色窗口找到网页端地址，复制到 Edge 或 Chrome 地址栏打开。"],
            ["网页提示无法访问", "等待半分钟后重试；仍失败时关闭所有黑色窗口，再重新双击启动文件。"],
            ["微信开发者工具没有打开", "手动打开开发者工具，选择导入项目，并选择整个 数字网格员 文件夹。"],
            ["开发者工具提示无权限", "确认使用的是参赛方已添加的评委微信号，然后退出并重新登录。"],
            ["模拟器能用 手机不能用", "按手机真机调试章节核对同一 Wi Fi 防火墙和后端 API 地址。"],
            ["端口不是 3100", "这是正常情况。程序会自动选择空闲端口，请始终使用本次窗口显示的地址。"],
        ],
        widths=[5.1, 10.4],
        center_cols=(),
    )

    add_heading(doc, "结束评审", 1)
    add_number(doc, "关闭浏览器和微信开发者工具。", start=True)
    add_number(doc, "关闭运行后端的黑色窗口，系统即停止。")
    add_number(doc, "如需再次体验，重新双击 启动数字网格员.bat。")
    add_body(doc, "评审结束后，参赛方可以在微信小程序成员管理中移除评委开发者权限。评委不需要进行任何账号解绑操作。")

    add_heading(doc, "评审检查清单", 1)
    add_check(doc, "网页端可以正常登录")
    add_check(doc, "可以提交诉求并看到处理过程")
    add_check(doc, "可以查看工单和政策来源")
    add_check(doc, "微信开发者工具可以编译小程序")
    add_check(doc, "手机可以打开并提交一条模拟诉求")

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build()
